from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="274E43")
ACCENT_FILL = PatternFill("solid", fgColor="E8F0EC")
WARN_FILL = PatternFill("solid", fgColor="FBE9E5")


def _write_table(ws, headers: list[str], rows: list[list], freeze: str = "A2") -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for index, column in enumerate(ws.columns, 1):
        width = max(len(str(cell.value or "")) for cell in column) + 2
        ws.column_dimensions[get_column_letter(index)].width = min(max(width, 12), 48)


def write_excel(payload: dict, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily_Summary"
    ws.append(["Synthetic Daily Abnormal Closure Report"])
    ws["A1"].font = Font(size=16, bold=True, color="274E43")
    ws.append([])
    ws.append(["Metric", "Value"])
    for cell in ws[3]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
    for label, value in (
        ("Anomalies", payload["totals"]["anomalies"]),
        ("Stores affected", payload["totals"]["stores"]),
        ("Abnormal minutes", payload["totals"]["abnormal_minutes"]),
        ("Estimated impacted sales (synthetic)", payload["totals"]["estimated_impacted_sales"]),
        ("Minimum abnormal minutes", payload["rules"]["min_abnormal_minutes"]),
        ("Impact factor", payload["rules"]["impact_factor"]),
    ):
        ws.append([label, value])
    ws["B7"].number_format = "¥#,##0.00"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.append([])
    ws.append([payload["notice"]])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=2)
    ws.cell(ws.max_row, 1).fill = WARN_FILL
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[ws.max_row].height = 42

    details_sheet = wb.create_sheet("Anomaly_Details")
    headers = ["date", "weekday", "channel", "store_id", "store_name", "region", "city", "supervisor", "planned_segments", "actual_segments", "abnormal_segments", "abnormal_minutes", "expected_hourly_sales", "estimated_impacted_sales", "orders", "gmv"]
    _write_table(details_sheet, headers, [[row[key] for key in headers] for row in payload["anomalies"]])
    for row in details_sheet.iter_rows(min_row=2, min_col=12, max_col=12):
        row[0].fill = WARN_FILL
    for row in details_sheet.iter_rows(min_row=2, min_col=14, max_col=14):
        row[0].number_format = "¥#,##0.00"

    for sheet_name, rows, keys in (
        ("Region_Summary", payload["region_summary"], ["region", "anomalies", "stores", "abnormal_minutes", "estimated_impacted_sales"]),
        ("Supervisor_Summary", payload["supervisor_summary"], ["region", "supervisor", "anomalies", "stores", "abnormal_minutes", "estimated_impacted_sales"]),
        ("Channel_Summary", payload["channel_summary"], ["channel", "anomalies", "stores", "abnormal_minutes", "estimated_impacted_sales"]),
    ):
        sheet = wb.create_sheet(sheet_name)
        _write_table(sheet, keys, [[row[key] for key in keys] for row in rows])
        for row in sheet.iter_rows(min_row=2, min_col=len(keys), max_col=len(keys)):
            row[0].number_format = "¥#,##0.00"

    method = wb.create_sheet("Methodology")
    method_rows = [
        ["Item", "Value"],
        ["Source mode", payload["source_mode"]],
        ["Impact formula", payload["rules"]["impact_formula"]],
        ["Minimum abnormal minutes", payload["rules"]["min_abnormal_minutes"]],
        ["Impact factor", payload["rules"]["impact_factor"]],
        ["Data classification", "Synthetic / simulated"],
        ["External notification", "Disabled by design"],
    ]
    for row in method_rows:
        method.append(row)
    for cell in method[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
    method.column_dimensions["A"].width = 32
    method.column_dimensions["B"].width = 72

    validation = wb.create_sheet("Input_Validation")
    _write_table(validation, ["check", "value"], [
        ["activity rows", payload["validation"]["activity_rows"]],
        ["matched rows", payload["validation"]["matched_rows"]],
        ["below threshold", payload["validation"]["below_threshold"]],
        ["missing business hours", len(payload["validation"]["missing_business_hours"])],
    ])
    wb.save(path)
    return path


def write_json(payload: dict, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def write_markdown(payload: dict, path: Path) -> Path:
    lines = [
        "# Synthetic Daily Abnormal Closure Report",
        "",
        f"> {payload['notice']}",
        "",
        "## Totals",
        f"- Anomalies: {payload['totals']['anomalies']}",
        f"- Stores affected: {payload['totals']['stores']}",
        f"- Abnormal minutes: {payload['totals']['abnormal_minutes']}",
        f"- Estimated impacted sales (synthetic): ¥{payload['totals']['estimated_impacted_sales']:,.2f}",
        "",
        "## Region summary",
        "| Region | Anomalies | Stores | Minutes | Estimated impact |",
        "| --- | ---: | ---: | ---: | ---: |",
    ]
    for row in payload["region_summary"]:
        lines.append(f"| {row['region']} | {row['anomalies']} | {row['stores']} | {row['abnormal_minutes']} | ¥{row['estimated_impacted_sales']:,.2f} |")
    lines += ["", "## Anomaly details", "| Date | Store | Channel | Gap | Minutes | Estimated impact |", "| --- | --- | --- | --- | ---: | ---: |"]
    for row in payload["anomalies"]:
        lines.append(f"| {row['date']} | {row['store_name']} | {row['channel']} | {row['abnormal_segments']} | {row['abnormal_minutes']} | ¥{row['estimated_impacted_sales']:,.2f} |")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path
