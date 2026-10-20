from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from closure_report.calculator import compute  # noqa: E402
from closure_report.config import Config, DEFAULTS  # noqa: E402
from closure_report.intervals import parse_segments, subtract_intervals, total_minutes  # noqa: E402
from closure_report.report import write_excel, write_json, write_markdown  # noqa: E402


class ClosureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        data = copy.deepcopy(DEFAULTS)
        data["app"]["archive_dir"] = str(Path(cls.temp.name) / "outputs")
        data["app"]["log_dir"] = str(Path(cls.temp.name) / "logs")
        cls.cfg = Config(root=ROOT, data=data)
        cls.payload = compute(cls.cfg)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_segment_subtraction(self):
        planned = parse_segments("10:00-22:00")
        actual = parse_segments("10:00-13:00;14:00-16:00;18:00-22:00")
        gaps = subtract_intervals(planned, actual)
        self.assertEqual(total_minutes(gaps), 180)
        self.assertEqual(len(gaps), 2)

    def test_synthetic_dataset_expected_counts(self):
        self.assertEqual(self.payload["validation"]["activity_rows"], 72)
        self.assertEqual(self.payload["totals"]["anomalies"], 6)
        self.assertEqual(self.payload["totals"]["stores"], 6)
        self.assertEqual(self.payload["totals"]["abnormal_minutes"], 1380)
        self.assertGreater(self.payload["totals"]["estimated_impacted_sales"], 0)
        self.assertTrue(self.payload["simulation"])

    def test_missing_hours_rejected_by_validation(self):
        bad = copy.deepcopy(DEFAULTS)
        bad["app"]["business_hours_path"] = "data/does-not-exist.csv"
        cfg = Config(root=ROOT, data=bad)
        with self.assertRaises(FileNotFoundError):
            compute(cfg)

    def test_excel_markdown_and_json_are_generated(self):
        path = Path(self.temp.name) / "outputs" / "report.xlsx"
        write_excel(self.payload, path)
        wb = load_workbook(path, read_only=True)
        self.assertIn("Anomaly_Details", wb.sheetnames)
        self.assertIn("Methodology", wb.sheetnames)
        self.assertEqual(wb["Anomaly_Details"].max_row, 7)
        wb.close()
        md = write_markdown(self.payload, Path(self.temp.name) / "summary.md")
        self.assertIn("synthetic", md.read_text(encoding="utf-8").lower())
        json_path = write_json(self.payload, Path(self.temp.name) / "summary.json")
        self.assertIn("estimated_impacted_sales", json_path.read_text(encoding="utf-8"))

    def test_public_mode_guard(self):
        data = copy.deepcopy(DEFAULTS)
        data["app"]["source_mode"] = "live"
        with self.assertRaises(ValueError):
            Config(root=ROOT, data=data).validate()


if __name__ == "__main__":
    unittest.main()
